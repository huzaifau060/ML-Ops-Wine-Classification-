import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Set standard 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Colors
PRIMARY = RGBColor(15, 23, 42)     # Slate 900
TEAL = RGBColor(13, 148, 136)       # Teal 600
DARK_BLUE = RGBColor(30, 58, 138)  # Blue 900
GRAY = RGBColor(100, 116, 139)     # Slate 500

# Helper to shade table cells
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

# Title & Header
p_inst = doc.add_paragraph()
r_inst = p_inst.add_run("FAST National University of Computer & Emerging Sciences")
r_inst.font.name = "Calibri"
r_inst.font.size = Pt(14)
r_inst.font.bold = True
r_inst.font.color.rgb = DARK_BLUE
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER

p_dept = doc.add_paragraph()
r_dept = p_dept.add_run("Department of Computer Science | Machine Learning Operations (MLOps)")
r_dept.font.name = "Calibri"
r_dept.font.size = Pt(11)
r_dept.font.color.rgb = GRAY
p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER

p_title = doc.add_paragraph()
r_title = p_title.add_run("Assignment 01: MLOps task with CI/CD, MLflow & Makefile Automation")
r_title.font.name = "Calibri"
r_title.font.size = Pt(18)
r_title.font.bold = True
r_title.font.color.rgb = PRIMARY
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_after = Pt(15)

# Student Info Box
info_table = doc.add_table(rows=3, cols=2)
info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
info_data = [
    ("Student Name:", "Huzaifa (Update with your full name)"),
    ("Roll Number:", "MLOps_A01_RollNumber (Update with your roll number)"),
    ("GitHub Repository:", "https://github.com/<your-username>/wine-mlops-pipeline"),
]
for i, (label, val) in enumerate(info_data):
    cell_lbl = info_table.cell(i, 0)
    cell_val = info_table.cell(i, 1)
    cell_lbl.width = Inches(2.2)
    cell_val.width = Inches(4.3)
    p0 = cell_lbl.paragraphs[0]
    r0 = p0.add_run(label)
    r0.font.bold = True
    r0.font.size = Pt(10)
    p1 = cell_val.paragraphs[0]
    r1 = p1.add_run(val)
    r1.font.size = Pt(10)
    set_cell_background(cell_lbl, "F1F5F9")
    set_cell_background(cell_val, "F8FAFC")

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# Section 1: Executive Overview & Table 1
h1 = doc.add_heading(level=1)
r_h1 = h1.add_run("1. Hyperparameter Search Results & Table 1")
r_h1.font.color.rgb = DARK_BLUE

p_intro = doc.add_paragraph(
    "A reproducible dual-classifier training pipeline was engineered on the Wine dataset "
    "(178 samples, 13 continuous chemical features, 3 cultivar classes) using a stratified "
    "80/20 train-test split (random_state=42). 5-fold stratified cross-validation was performed "
    "across 3 distinct hyperparameter configurations for RandomForestClassifier and "
    "3 configurations for GradientBoostingClassifier."
)
p_intro.paragraph_format.space_after = Pt(8)

# Table 1
table1 = doc.add_table(rows=7, cols=7)
table1.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Model Family", "Config Name", "Key Hyperparameters", "Val Macro F1", "Val Accuracy", "Val Log Loss", "Status"]
rows_data = [
    ["RandomForest", "rf_config_1", "n_est=50, depth=3, split=2", "0.9665", "0.9653", "0.2055", "Candidate"],
    ["RandomForest", "rf_config_2", "n_est=100, depth=5, split=2", "0.9789", "0.9791", "0.1630", "CHAMPION"],
    ["RandomForest", "rf_config_3", "n_est=150, depth=None, split=4", "0.9789", "0.9791", "0.1688", "Candidate"],
    ["GradientBoosting", "gbm_config_1", "n_est=50, lr=0.05, depth=3", "0.9053", "0.9022", "0.2295", "Candidate"],
    ["GradientBoosting", "gbm_config_2", "n_est=100, lr=0.1, depth=3", "0.9259", "0.9236", "0.4281", "Candidate"],
    ["GradientBoosting", "gbm_config_3", "n_est=150, lr=0.2, depth=4", "0.9197", "0.9160", "0.5264", "Candidate"],
]

for col_idx, h in enumerate(headers):
    cell = table1.cell(0, col_idx)
    p = cell.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(255, 255, 255)
    set_cell_background(cell, "1E293B")

for row_idx, rdata in enumerate(rows_data, start=1):
    for col_idx, val in enumerate(rdata):
        cell = table1.cell(row_idx, col_idx)
        p = cell.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(8.5)
        if rdata[6] == "CHAMPION":
            set_cell_background(cell, "DCFCE7")
            r.font.bold = True
            r.font.color.rgb = RGBColor(22, 101, 52)
        elif row_idx % 2 == 0:
            set_cell_background(cell, "F8FAFC")
        else:
            set_cell_background(cell, "FFFFFF")

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# Add Figure 1
doc.add_picture("report_materials/figures/figure1_hyperparameter_metrics.png", width=Inches(6.2))
p_fig1 = doc.add_paragraph("Figure 1: Hyperparameter validation comparison across 5-fold cross-validation.")
p_fig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_fig1.runs[0].font.size = Pt(8.5)
p_fig1.runs[0].font.italic = True
p_fig1.paragraph_format.space_after = Pt(15)

# Section 2: MLflow UI Visualizations
doc.add_page_break()
h2 = doc.add_heading(level=1)
r_h2 = h2.add_run("2. MLflow Tracking & Model Registry Visualizations")
r_h2.font.color.rgb = DARK_BLUE

doc.add_paragraph(
    "MLflow was configured with a SQLite backend store (sqlite:///mlflow.db). "
    "All 6 candidate configurations logged parameters, metrics, model signatures, "
    "and input examples under the experiment 'Wine-Cultivar-Classification'. "
    "The winning model (rf_config_2) was registered as 'WineClassifier' with alias 'champion'."
)

doc.add_picture("report_materials/figures/figure2_mlflow_experiment_overview.png", width=Inches(6.2))
p_fig2 = doc.add_paragraph("Figure 2: MLflow experiment overview showing logged runs, hyperparameters, and metrics.")
p_fig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_fig2.runs[0].font.size = Pt(8.5)
p_fig2.runs[0].font.italic = True
p_fig2.paragraph_format.space_after = Pt(12)

doc.add_picture("report_materials/figures/figure3_mlflow_model_registry.png", width=Inches(6.2))
p_fig3 = doc.add_paragraph("Figure 3: MLflow Model Registry displaying 'WineClassifier' promoted with '@champion' alias.")
p_fig3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_fig3.runs[0].font.size = Pt(8.5)
p_fig3.runs[0].font.italic = True
p_fig3.paragraph_format.space_after = Pt(15)

# Section 3: CI/CD Pipeline & Quality Gate Evidence
h3 = doc.add_heading(level=1)
r_h3 = h3.add_run("3. Automated CI/CD Pipeline & MLOps Quality Gate")
r_h3.font.color.rgb = DARK_BLUE

doc.add_paragraph(
    "The GitHub Actions workflow (.github/workflows/ci.yml) triggers on every push and pull request to 'main'. "
    "It executes on ubuntu-latest with Python 3.10 and enforces three automated MLOps Quality Gates in tests/test_model_gate.py:"
)

gates_bullet = [
    ("Metric Threshold Gate: ", "Validation Macro F1 >= 0.88 (Champion achieved 0.9789 - PASSED)"),
    ("Inference Latency Gate: ", "Batch inference time <= 30 ms (Champion achieved < 1.5 ms - PASSED)"),
    ("Output Schema Integrity Gate: ", "Model predictions strictly restricted to class indices {0, 1, 2} - PASSED"),
]
for g_title, g_desc in gates_bullet:
    p = doc.add_paragraph(style='List Bullet')
    r_t = p.add_run(g_title)
    r_t.font.bold = True
    p.add_run(g_desc)

doc.add_picture("report_materials/figures/figure5_github_actions_ci_flow.png", width=Inches(6.2))
p_fig5 = doc.add_paragraph("Figure 4: GitHub Actions automated CI execution pipeline stages.")
p_fig5.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_fig5.runs[0].font.size = Pt(8.5)
p_fig5.runs[0].font.italic = True
p_fig5.paragraph_format.space_after = Pt(15)

# Section 4: Git Collaboration & Merge Conflict Resolution
doc.add_page_break()
h4 = doc.add_heading(level=1)
r_h4 = h4.add_run("4. Git Collaboration & Conflict Resolution (Milestone 5)")
r_h4.font.color.rgb = DARK_BLUE

doc.add_paragraph(
    "To demonstrate version control proficiency, development followed a clean feature-branch "
    "workflow ('feature/mlflow-tracking'). An engineered merge conflict was created on branch "
    "'conflict-simulation' modifying the F1 threshold in src/config.py concurrently with a conflicting edit on 'main'."
)

doc.add_picture("report_materials/figures/figure4_git_commit_tree.png", width=Inches(6.2))
p_fig4 = doc.add_paragraph("Figure 5: Full Git commit tree demonstrating feature branch PR and resolved merge conflict.")
p_fig4.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_fig4.runs[0].font.size = Pt(8.5)
p_fig4.runs[0].font.italic = True
p_fig4.paragraph_format.space_after = Pt(10)

p_code_header = doc.add_paragraph()
r_ch = p_code_header.add_run("Terminal Output of Conflict Simulation & Manual Resolution:")
r_ch.font.bold = True
r_ch.font.size = Pt(9.5)

terminal_text = (
    "$ git merge conflict-simulation\n"
    "Auto-merging src/config.py\n"
    "CONFLICT (content): Merge conflict in src/config.py\n"
    "Automatic merge failed; fix conflicts and then commit the result.\n\n"
    "$ cat src/config.py\n"
    "<<<<<<< HEAD\n"
    "METRIC_THRESHOLD_F1 = 0.88  # Mandatory assignment baseline specification\n"
    "=======\n"
    "METRIC_THRESHOLD_F1 = 0.90  # Quality gate proposed threshold from experimentation\n"
    ">>>>>>> conflict-simulation\n\n"
    "$ # Resolution: Reconciled experimental suggestion with mandatory baseline\n"
    "$ git add src/config.py\n"
    "$ git commit -m \"Merge branch 'conflict-simulation' and resolve merge conflict in src/config.py\"\n"
    "[main 985f384] Merge branch 'conflict-simulation' and resolve merge conflict in src/config.py"
)
p_term = doc.add_paragraph()
r_term = p_term.add_run(terminal_text)
r_term.font.name = "Consolas"
r_term.font.size = Pt(8)
r_term.font.color.rgb = PRIMARY
p_term.paragraph_format.left_indent = Inches(0.2)
p_term.paragraph_format.space_after = Pt(15)

# Section 5: Analytical Response (Max 150 words)
h5 = doc.add_heading(level=1)
r_h5 = h5.add_run("5. Analytical Response (Max 150 words)")
r_h5.font.color.rgb = DARK_BLUE

analysis_text = (
    "A comparative analysis between Random Forest and Gradient Boosting demonstrates distinct "
    "performance dynamics on the chemical cultivar dataset. Random Forest achieved superior generalization, "
    "with rf_config_2 reaching a peak validation Macro F1 of 0.9789 and the lowest validation Log Loss of 0.1630. "
    "Its ensemble averaging effectively mitigated overfitting across the 13 continuous chemical features. "
    "In contrast, Gradient Boosting attained a maximum Macro F1 of 0.9259 under gbm_config_2. While sequential "
    "boosting refined residual errors, higher learning rates and depths led to elevated validation log loss "
    "(0.5264 in gbm_config_3), indicating sensitivity to the relatively small sample size (142 training instances). "
    "Random Forest (rf_config_2) proved more robust and was selected as the production champion."
)

p_anal = doc.add_paragraph(analysis_text)
p_anal.paragraph_format.line_spacing = 1.15
p_anal.paragraph_format.space_after = Pt(10)

word_count = len(analysis_text.split())
p_cnt = doc.add_paragraph(f"(Word Count: {word_count} words / 150 maximum permitted)")
p_cnt.runs[0].font.size = Pt(8.5)
p_cnt.runs[0].font.italic = True
p_cnt.runs[0].font.color.rgb = GRAY

doc.save("report_materials/MLOps_A01_Report.docx")
print("MLOps_A01_Report.docx successfully created!")
